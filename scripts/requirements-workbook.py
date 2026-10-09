#!/usr/bin/env python3
"""Refresh docs/HMS_R1_Requirements.xlsx from docs/requirements.csv (the single register, wave 1 item 3).

Rewrites the "R1 Requirements" sheet and the counts on "README". Other sheets are kept as they are.
Needs openpyxl (pip install openpyxl). Run: python3 scripts/requirements-workbook.py
"""
import collections, csv, datetime
import openpyxl
from openpyxl.styles import Alignment, Font

BOOK = "docs/HMS_R1_Requirements.xlsx"
rows = list(csv.reader(open("docs/requirements.csv", newline="")))
head, body = rows[0], rows[1:]
wb = openpyxl.load_workbook(BOOK)

ws = wb["R1 Requirements"]
widths = {c: ws.column_dimensions[c].width for c in ws.column_dimensions}
ws.delete_rows(1, ws.max_row)
ws.append(head)
for r in body:
    ws.append(r)
for cell in ws[1]:
    cell.font = Font(bold=True)
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for c, w in widths.items():
    ws.column_dimensions[c].width = w
ws.freeze_panes = "B2"

rd = wb["README"]
status = collections.Counter(r[head.index("Status")] for r in body)
for row in rd.iter_rows():
    if row[0].value and str(row[0].value).startswith("Version"):
        row[0].value = f"Generated from docs/requirements.csv on {datetime.date.today():%d %b %Y}. Status describes the Zikora app."
start = next(i for i, row in enumerate(rd.iter_rows(), 1) if row[0].value == "Counts")
rd.delete_rows(start + 1, rd.max_row)
rd.cell(start + 1, 1, "Total requirements"); rd.cell(start + 1, 2, len(body))
for n, (k, v) in enumerate(sorted(status.items()), start + 2):
    rd.cell(n, 1, k); rd.cell(n, 2, v)
wb.save(BOOK)
print(f"{BOOK}: {len(body)} rows; " + ", ".join(f"{k} {v}" for k, v in sorted(status.items())))
